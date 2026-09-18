import csv
import openpyxl
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from auto_gen_package.app_init import DIR_OUTPUT, DIR_LOG
from auto_gen_package import app_init
from auto_gen_package.format import lookup, group_name_transform, disk_text_format, get_net_out_config

lst_online = []
#模板制作
def make_template(template:Workbook,master:Worksheet,scmt:str,sync_date:str):
    tp_job = template["迁移任务"]
    tp_hci_info = template["虚拟平台信息"]
    need_del = []

    ##填充数据，列有变动全在这里修改
    for row in tp_job.iter_rows(min_row=5):

        #任务分组[0]
        row[0].value = group_name_transform(lookup(master,row[1].value,app_init.MASTER_HEADER["桌面ID"],app_init.MASTER_HEADER["迁移日"]))
        if row[0].value is not None:
            lst_online.append(row[1].value)

        #删除无任务的记录
        if row[0].value is None:
            need_del.append(row[0].row)

        #目标机名称[3]
        row[3].value = lookup(master,row[1].value,app_init.MASTER_HEADER["桌面ID"],app_init.MASTER_HEADER["新桌面名称"])

        #目标虚拟平台[4]
        hci_ip = lookup(master,row[1].value,app_init.MASTER_HEADER["桌面ID"],app_init.MASTER_HEADER["目标HCI"])
        if hci_ip is not None:
            row[4].value = f"Sangfor HCI（{hci_ip}）"

        #虚拟机分组[5]
        row[5].value ="迁移服务器"+scmt

        #存储位置[7]
        row[7].value = lookup(master,row[1].value,app_init.MASTER_HEADER["桌面ID"],app_init.MASTER_HEADER["存储位置"])

        #同步磁盘[10]

        row[10].value = disk_text_format(row[10].value)
        #同步开始时间[11]
        row[11].value=sync_date

        #同步周期[12]
        row[12].value = "不启用"
        #网络出口配置[13]
        row[13].value = get_net_out_config(master,tp_hci_info,row[13].value,row[1].value)

        #网关配置[14]
        row[14].value = None
        #DNS配置[15]
        row[15].value = None

    #倒删
    for r_no in sorted(need_del, reverse=True):
        tp_job.delete_rows(r_no)
    return template

#csv
def make_script_ips(master:Worksheet):

    #表头
    header = ["目标机名","用户名","是否双屏","是否4k","资源名","角色"]

    #获取数据，渲染数据模型
    rows = []
    for row in master.iter_rows(min_row=2,values_only=True):
        #检查在清单里是否满足迁移要求
        if group_name_transform(row[app_init.MASTER_HEADER["迁移日"]]) is None:
            continue
        lst = [row[app_init.MASTER_HEADER["新桌面名称"]],row[app_init.MASTER_HEADER["使用人工号"]],1,1,row[app_init.MASTER_HEADER["vdc资源名称"]],row[app_init.MASTER_HEADER["vdc角色"]]]
        rows.append(lst)
    with open(DIR_OUTPUT / "output-脚本导入ips.csv",'w',newline='',encoding='utf-8-sig') as csvfile:

        # writer = csv.DictWriter(csvfile,fieldnames=fieldnames)
        # writer.writeheader()
        wtr = csv.writer(csvfile)
        wtr.writerow(header)
        for row in rows:
            wtr.writerow(row)
#csv
def make_html_info(master:Worksheet):
    #初始化表头
    header = [
        ["#每行前面的'#'代表该行是注释行，而不是需导入的数据",None,None,None,None,None],
        ["#请参考下面的示例添加需导入的信息，带*的字段为必填项，请不要删除列及修改列的顺序",None,None,None,None,None],
        ["#如果只需要导入用户名，只需要填写用户名",None,None,None,None,None],
        ["#如果只需要导入虚拟机IP，只需要填写虚拟机对应的IP、子网掩码、默认网关、首选DNS服务器",None,None,None,None,None],
        ["#如果需要导入IP和用户，则需要填写表格内所有信息。",None,None,None,None,None],
        ["#数据导入失败，请下载表格查看详细的失败信息",None,None,None,None,None],
        ["#虚拟机名称(*)","虚拟机IP","子网掩码","默认网关", "首选DNS服务器", "用户名"]
    ]
    #写入表头
    with open(DIR_OUTPUT /  "output-虚拟机关联用户.csv",'w',newline='',encoding='utf-8-sig') as csvfile:
        wtr = csv.writer(csvfile)
        for row in header:
            wtr.writerow(row)

        #初始化数据
        rows = []
        for row in master.iter_rows(min_row=2,values_only=True):
            #检查在清单里是否满足迁移要求
            if group_name_transform(row[app_init.MASTER_HEADER["迁移日"]]) is None:
                continue
            lst = [row[app_init.MASTER_HEADER["新桌面名称"]],None,None,None,None,row[app_init.MASTER_HEADER["使用人工号"]]]
            rows.append(lst)

        #写入数据
        for row in rows:
            wtr.writerow(row)


def make_dest_monitor_info(master:Worksheet):
    #初始化表
    wb = openpyxl.Workbook()
    wb_sheet = wb.active
    headers = ["工号","姓名","IP","分组","注册表值"]
    wb_sheet.append(headers)

    #填充数据
    for row in master.iter_rows(min_row=2,values_only=True):
        #检查在清单里是否满足迁移要求
        if group_name_transform(row[app_init.MASTER_HEADER["迁移日"]]) is None:
            continue
        lst = [row[app_init.MASTER_HEADER["使用人工号"]],row[app_init.MASTER_HEADER["使用人姓名"]],row[app_init.MASTER_HEADER["新桌面IP"]],row[app_init.MASTER_HEADER["迁移服务器"]],None]
        wb_sheet.append(lst)

    return wb

def make_not_online(lst_ol:list,master:Worksheet):
    #输出不在线情况
    wb = openpyxl.Workbook()
    ws = wb.active
    header = ["不在线虚拟机","迁移服务器"]
    ws.append(header)
    for row in master.iter_rows(min_row=2):
        #检查在清单里是否满足迁移要求
        if group_name_transform(row[app_init.MASTER_HEADER["迁移日"]].value) is None:
            continue
        if row[app_init.MASTER_HEADER["桌面ID"]].value not in lst_ol:
            ws.append([row[app_init.MASTER_HEADER["桌面ID"]].value,row[app_init.MASTER_HEADER["迁移服务器"]].value])
    wb.save(DIR_LOG / "不在线情况统计表.xlsx")

def make_total_template_logs(total:list[Workbook]):
    wb = openpyxl.Workbook()
    ws = wb.active
    header = [
        "任务分组",
        "源机名称*",
        "源机IP",
        "目标机名称",
        "目标虚拟平台*",
        "虚拟机分组*",
        "运行位置",
        "存储位置*",
        "目标CPU核数",
        "目标内存(GB)",
        "同步磁盘",
        "同步开始时间",
        "同步周期",
        "网络出口配置",
        "网关配置",
        "DNS配置"
    ]
    ws.append(header)
    for wb_it in total:
        ws_it = wb_it.active
        for row in ws_it.iter_rows(min_row=5):
            row_values = [cell.value for cell in row]
            ws.append(row_values)

    wb.save(DIR_LOG / "任务汇总.xlsx")
    wb.close()
